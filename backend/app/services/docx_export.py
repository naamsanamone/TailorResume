import io
import logging
from typing import List, Dict, Any
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

logger = logging.getLogger(__name__)

async def generate_docx(resume_sections: List[Dict[str, Any]]) -> bytes:
    """Generate DOCX bytes from resume sections."""
    doc = Document()
    
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)
        
    for sec in resume_sections:
        sec_type = sec.get("type", "").lower()
        
        if sec_type == "header":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(sec.get("fullName", ""))
            r.bold = True
            r.font.size = Pt(16)
            
            contact_parts = []
            if sec.get("email"): contact_parts.append(sec.get("email"))
            if sec.get("phone"): contact_parts.append(sec.get("phone"))
            if sec.get("location"): contact_parts.append(sec.get("location"))
            if sec.get("linkedin"): contact_parts.append(sec.get("linkedin"))
            
            p2 = doc.add_paragraph(" | ".join(contact_parts))
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
        else:
            doc.add_paragraph()
            p_title = doc.add_paragraph()
            r_title = p_title.add_run(sec.get("name", sec_type.title()).upper())
            r_title.bold = True
            
            if sec_type == "summary":
                doc.add_paragraph(sec.get("text", ""))
                
            elif sec_type == "skills":
                cats = sec.get("categories", {})
                if isinstance(cats, dict) and cats:
                    for cat_name, cat_skills in cats.items():
                        p = doc.add_paragraph()
                        r_cat = p.add_run(f"{cat_name}: ")
                        r_cat.bold = True
                        p.add_run(cat_skills if isinstance(cat_skills, str) else ", ".join(cat_skills))
                elif sec.get("items"):
                    doc.add_paragraph(", ".join(sec.get("items", [])))
                    
            elif sec_type in ["experience", "education", "projects"]:
                for entry in sec.get("entries", []):
                    title = entry.get("title", "") or entry.get("degree", "") or entry.get("name", "")
                    org = entry.get("company", "") or entry.get("institution", "")
                    date = entry.get("duration", "") or entry.get("date", "") or entry.get("year", "")
                    loc = entry.get("location", "")
                    
                    p1 = doc.add_paragraph()
                    r_title_entry = p1.add_run(title)
                    r_title_entry.bold = True
                    p1.add_run(f" \t {date}")
                    
                    p2 = doc.add_paragraph()
                    r_org = p2.add_run(org)
                    r_org.italic = True
                    p2.add_run(f" \t {loc}")
                    
                    for bullet in entry.get("bullets", []):
                        doc.add_paragraph(bullet, style='List Bullet')
                        
            elif sec_type == "list":
                for item in sec.get("items", []):
                    doc.add_paragraph(item, style='List Bullet')

    file_stream = io.BytesIO()
    doc.save(file_stream)
    return file_stream.getvalue()
