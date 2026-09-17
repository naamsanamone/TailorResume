import io
import logging
import fitz
from docx import Document
from typing import List, Dict, Any
from app.services.llm_client import get_llm_client
from app.prompts.templates import PARSE_RESUME_PROMPT

logger = logging.getLogger(__name__)

async def parse_resume_pdf(file_bytes: bytes) -> str:
    """Extract text from PDF bytes using PyMuPDF."""
    text = ""
    try:
        pdf_document = fitz.open(stream=file_bytes, filetype="pdf")
        for page_num in range(pdf_document.page_count):
            page = pdf_document.load_page(page_num)
            text += page.get_text() + "\n\n"
        pdf_document.close()
    except Exception as e:
        logger.error(f"Error parsing PDF: {str(e)}")
        raise ValueError("Failed to parse PDF file.")
    return text

async def parse_resume_docx(file_bytes: bytes) -> str:
    """Extract text from DOCX bytes using python-docx."""
    text = ""
    try:
        doc = Document(io.BytesIO(file_bytes))
        for para in doc.paragraphs:
            if para.text.strip():
                text += para.text + "\n"
        for table in doc.tables:
            for row in table.rows:
                row_data = []
                for cell in row.cells:
                    if cell.text.strip():
                        row_data.append(cell.text.strip())
                if row_data:
                    text += " | ".join(row_data) + "\n"
    except Exception as e:
        logger.error(f"Error parsing DOCX: {str(e)}")
        raise ValueError("Failed to parse DOCX file.")
    return text

async def structure_resume(raw_text: str) -> List[Dict[str, Any]]:
    """Convert raw text into structured ResumeSection[] JSON using LLM."""
    client = get_llm_client(for_content=False)
    system_message = "You are an expert resume parser. Extract the resume information strictly according to the provided JSON schema."
    prompt = PARSE_RESUME_PROMPT.format(raw_text=raw_text)
    
    try:
        result = await client.complete_json(
            prompt=prompt,
            system_message=system_message,
            temperature=0.1
        )
        return result.get("sections", [])
    except Exception as e:
        logger.error(f"Error structuring resume: {str(e)}")
        raise ValueError("Failed to structure resume content.")

async def parse_resume_file(filename: str, file_bytes: bytes) -> List[Dict[str, Any]]:
    """Dispatch to pdf or docx parser based on extension, then structure."""
    ext = filename.split(".")[-1].lower()
    
    if ext == "pdf":
        raw_text = await parse_resume_pdf(file_bytes)
    elif ext in ["doc", "docx"]:
        raw_text = await parse_resume_docx(file_bytes)
    else:
        raise ValueError(f"Unsupported file extension: {ext}")
        
    if not raw_text.strip():
        raise ValueError("No text could be extracted from the file.")
        
    return await structure_resume(raw_text)
